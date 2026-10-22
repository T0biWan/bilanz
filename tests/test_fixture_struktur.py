"""Prüft, dass f1_ok.xlsx genauso aufgebaut ist wie ein echter Finanzguru-Export.

Läuft nur, wenn ein echter Export angegeben ist (sonst übersprungen):
    set BILANZ_EXPORT=daten\\export.xlsx
    pytest tests/test_fixture_struktur.py
Gelesen werden nur Kopfzeile und Zelltypen – der Export bleibt lokal.
"""

import os
import warnings
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

ECHTER_EXPORT = os.environ.get("BILANZ_EXPORT")
FIXTURES = Path(__file__).parent / "fixtures"


def art(wert):
    if isinstance(wert, datetime):
        return "datum"
    if isinstance(wert, (int, float)):  # int/float gelten als gleich
        return "zahl"
    return "text"


def struktur(pfad):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(pfad, read_only=True)
    zeilen = wb.active.iter_rows(values_only=True)
    kopf = list(next(zeilen))
    arten = {spalte: set() for spalte in kopf}
    for z in zeilen:
        for spalte, wert in zip(kopf, z):
            if wert is not None and wert != "":
                arten[spalte].add(art(wert))
    wb.close()
    return kopf, arten


@pytest.mark.skipif(not ECHTER_EXPORT, reason="BILANZ_EXPORT nicht gesetzt")
def test_f1_ok_ist_aufgebaut_wie_echter_export():
    kopf_echt, arten_echt = struktur(ECHTER_EXPORT)
    kopf_fix, arten_fix = struktur(FIXTURES / "f1_ok.xlsx")
    assert kopf_fix == kopf_echt
    for spalte in kopf_fix:
        assert arten_fix[spalte] <= arten_echt[spalte], spalte
