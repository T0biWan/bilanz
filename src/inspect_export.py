"""bilanz – Schritt 1: Finanzguru-Export prüfen.

- Zeigt pro Spalte Füllgrad und Python-Typen (keine Daten, kannst du mir schicken).
- Schreibt alle Gegenparteien nach keep.txt (bleibt bei dir).

Aufruf: python inspect_export.py <export.xlsx>
"""

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = [
    "Buchungstag", "Referenzkonto", "Name Referenzkonto", "Betrag", "Kontostand",
    "Waehrung", "Beguenstigter/Auftraggeber", "IBAN Beguenstigter/Auftraggeber",
    "Verwendungszweck", "E-Ref", "Mandatsreferenz", "Glaeubiger-ID",
    "Analyse-Hauptkategorie", "Analyse-Unterkategorie", "Analyse-Vertrag",
    "Analyse-Vertragsturnus", "Analyse-Vertrags-ID", "Analyse-Umbuchung",
    "Analyse-Vom frei verfuegbaren Einkommen ausgeschlossen", "Analyse-Umsatzart",
    "Analyse-Betrag", "Analyse-Woche", "Analyse-Monat", "Analyse-Quartal",
    "Analyse-Jahr", "Buchungs-ID", "Referenz-Original-ID", "Split-Typ", "Tags", "Notiz",
]
PARTY = "Beguenstigter/Auftraggeber"
KEEP_FILE = Path("keep.txt")
KEEP_HEADER = """\
# Alle Gegenparteien, sortiert nach Häufigkeit.
# LÖSCHE alle Zeilen mit Privatpersonen oder sonst Sensiblem.
# Was hier stehen bleibt, wird beim Anonymisieren NICHT ersetzt.
# Alles andere (auch neue Namen in späteren Exporten) wird pseudonymisiert.
"""


def is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def main(path: Path) -> None:
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    print(f"Sheets: {wb.sheetnames} (lese: {ws.title})")

    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    data = [r for r in rows if not all(is_empty(v) for v in r)]
    wb.close()

    missing = [c for c in EXPECTED if c not in header]
    extra = [c for c in header if c not in EXPECTED]
    print(f"Buchungen: {len(data)}")
    print(f"Fehlende Spalten: {missing or '-'}")
    print(f"Unerwartete Spalten: {extra or '-'}")
    print()
    print(f"{'Spalte':<56} {'gefüllt':>8}  Typen")
    for i, col in enumerate(header):
        values = [r[i] for r in data if not is_empty(r[i])]
        types = sorted({type(v).__name__ for v in values})
        print(f"{str(col):<56} {len(values):>8}  {', '.join(types)}")

    i = header.index(PARTY)
    counts = Counter(str(r[i]).strip() for r in data if not is_empty(r[i]))
    names = [name for name, _ in counts.most_common()]
    KEEP_FILE.write_text(KEEP_HEADER + "\n".join(names) + "\n", encoding="utf-8")
    print(f"\n{len(names)} Gegenparteien -> {KEEP_FILE}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Aufruf: python inspect_export.py <export.xlsx>")
    main(Path(sys.argv[1]))
