"""F1 – Finanzguru-Export einlesen.

Spezifikation: docs/features/F1-import.md
"""

import warnings
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

PFLICHTSPALTEN = [
    "Buchungstag", "Referenzkonto", "Name Referenzkonto", "Betrag", "Kontostand", "Waehrung",
    "Beguenstigter/Auftraggeber", "IBAN Beguenstigter/Auftraggeber",
    "Verwendungszweck", "Buchungs-ID", "Split-Typ",
]
SPLIT_TEILE = {"Teilbuchung", "Restbetrag"}
CENT = Decimal("0.01")


@dataclass(frozen=True)
class Buchung:
    zeile: int                # Excel-Zeile im Export (für Meldungen)
    id: str                   # Buchungs-ID von Finanzguru
    konto: str                # Referenzkonto (IBAN oder ID)
    konto_name: str | None    # Kontoname aus Finanzguru ("Name Referenzkonto"), nur Vorschlag für F2
    datum: date              # nur der Tag, Uhrzeit fällt weg
    betrag: Decimal           # exakt, immer 2 Nachkommastellen
    gegenpartei: str | None
    gegen_iban: str | None
    zweck: str | None


@dataclass(frozen=True)
class Warnung:
    konto: str
    zeile: int
    erwartet: Decimal         # Kontostand laut Vorgänger + Betrag
    gefunden: Decimal         # Kontostand laut Export


@dataclass(frozen=True)
class ImportErgebnis:
    buchungen: list[Buchung]
    warnungen: list[Warnung]


class ImportFehler(Exception):
    """Abbruch: der Export kann nicht verlässlich gelesen werden."""

    def __init__(self, grund: str, meldung: str, zeilen: list[int] | None = None):
        super().__init__(meldung)
        self.grund = grund
        self.zeilen = zeilen or []


def load_export(pfad: str | Path) -> ImportErgebnis:
    zeilen = [(n, z) for n, z in _lies_zeilen(Path(pfad)) if z["Split-Typ"] not in SPLIT_TEILE]
    eintraege = [_lies_buchung(n, z) for n, z in zeilen]
    buchungen = [b for b, _ in eintraege]
    _pruefe_ids(buchungen)
    return ImportErgebnis(buchungen, _pruefe_kontostaende(eintraege))


# --- Lesen -------------------------------------------------------------------

def _lies_zeilen(pfad: Path) -> list[tuple[int, dict]]:
    """Alle nicht-leeren Datenzeilen als (Excel-Zeile, {Spalte: Wert})."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        wb = load_workbook(pfad, read_only=True)
    try:
        werte_je_zeile = wb.active.iter_rows(values_only=True)
        kopf = next(werte_je_zeile, None) or ()
        fehlend = [s for s in PFLICHTSPALTEN if s not in kopf]
        if fehlend:
            raise ImportFehler("fehlende_spalte", f"Pflichtspalte fehlt: {', '.join(fehlend)}")
        return [(n, dict(zip(kopf, werte)))
                for n, werte in enumerate(werte_je_zeile, start=2)
                if not all(_leer(w) for w in werte)]
    finally:
        wb.close()


def _lies_buchung(n: int, z: dict) -> tuple[Buchung, Decimal]:
    """Eine Zeile -> (Buchung, Kontostand). Bricht bei ungültigen Pflichtfeldern ab."""

    def pflicht(spalte):
        wert = z.get(spalte)
        if _leer(wert):
            raise ImportFehler("leeres_feld", f"Zeile {n}: {spalte} ist leer", [n])
        return wert

    def ungueltig(spalte, grund):
        return ImportFehler("ungueltiges_feld", f"Zeile {n}: {spalte} {grund}: {z.get(spalte)!r}", [n])

    waehrung = z.get("Waehrung")
    if waehrung != "EUR":
        raise ImportFehler("waehrung", f"Zeile {n}: Währung {waehrung!r}, unterstützt wird nur EUR", [n])

    tag = pflicht("Buchungstag")
    if not isinstance(tag, datetime):
        raise ungueltig("Buchungstag", "ist kein Datum")

    for spalte in ("Buchungs-ID", "Referenzkonto"):
        if not isinstance(pflicht(spalte), str):
            raise ungueltig(spalte, "ist kein Text")

    def betrag(spalte):
        wert = pflicht(spalte)
        if isinstance(wert, bool) or not isinstance(wert, (int, float)):
            raise ungueltig(spalte, "ist keine Zahl")
        exakt = Decimal(str(wert))  # str() statt float: -0.1 statt -0.1000000000000000055
        if not exakt.is_finite() or exakt.as_tuple().exponent < -2:
            raise ungueltig(spalte, "hat mehr als 2 Nachkommastellen")
        return exakt.quantize(CENT)  # nur Stellen auffüllen, nie runden

    buchung = Buchung(
        zeile=n,
        id=z["Buchungs-ID"],
        konto=z["Referenzkonto"],
        konto_name=_text(z["Name Referenzkonto"]),
        datum=tag.date(),
        betrag=betrag("Betrag"),
        gegenpartei=_text(z["Beguenstigter/Auftraggeber"]),
        gegen_iban=_text(z["IBAN Beguenstigter/Auftraggeber"]),
        zweck=_text(z["Verwendungszweck"]),
    )
    return buchung, betrag("Kontostand")


def _leer(wert) -> bool:
    return wert is None or (isinstance(wert, str) and not wert.strip())


def _text(wert) -> str | None:
    return None if _leer(wert) else str(wert)


# --- Prüfen ------------------------------------------------------------------

def _pruefe_ids(buchungen: list[Buchung]) -> None:
    zeilen_je_id: dict[str, list[int]] = {}
    for b in buchungen:
        zeilen_je_id.setdefault(b.id, []).append(b.zeile)
    doppelt = sorted(z for zeilen in zeilen_je_id.values() if len(zeilen) > 1 for z in zeilen)
    if doppelt:
        raise ImportFehler("doppelte_id", f"Buchungs-ID mehrfach vergeben in Zeilen {doppelt}", doppelt)


def _pruefe_kontostaende(eintraege: list[tuple[Buchung, Decimal]]) -> list[Warnung]:
    """Pro Konto muss gelten: Kontostand = Kontostand der nächstälteren Zeile + Betrag.

    Der Export ist absteigend sortiert, die nächstältere Zeile steht also darunter.
    """
    je_konto: dict[str, list[tuple[Buchung, Decimal]]] = {}
    for b, stand in eintraege:
        je_konto.setdefault(b.konto, []).append((b, stand))

    warnungen = []
    for liste in je_konto.values():
        for (b, stand), (_, stand_davor) in zip(liste, liste[1:]):
            erwartet = stand_davor + b.betrag
            if stand != erwartet:
                warnungen.append(Warnung(konto=b.konto, zeile=b.zeile, erwartet=erwartet, gefunden=stand))
    return sorted(warnungen, key=lambda w: w.zeile)
