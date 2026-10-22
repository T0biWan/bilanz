"""bilanz – Schritt 2: Finanzguru-Export lokal anonymisieren.

Aufruf: python anonymize.py <export.xlsx> [keep.txt] [redact.txt]
Ergebnis: <export>_anon.xlsx – gleiche Struktur, gleiche Zelltypen.

keep.txt      Gegenparteien/Kontonamen, die im Klartext bleiben (eine pro Zeile).
redact.txt    Optional: Begriffe, die überall zu XXX werden (dein Name, Adresse …).
anon_key.txt  Geheimer Schlüssel, wird beim ersten Lauf erzeugt. NICHT teilen.
              Gleicher Schlüssel => gleiche Pseudonyme über alle Exporte hinweg.

Unverändert bleiben: Beträge, Kontostände, Daten, Buchungs-IDs, Finanzguru-Kategorien.
"""

import hashlib
import hmac
import re
import secrets
import sys
import warnings
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

warnings.filterwarnings("ignore", message="Workbook contains no default style")

KEY_FILE = Path("anon_key.txt")

PARTY = "Beguenstigter/Auftraggeber"
ACCOUNT_NAME = "Name Referenzkonto"
IBAN_COLS = ["Referenzkonto", "IBAN Beguenstigter/Auftraggeber"]
ID_COLS = {"E-Ref": "EREF", "Mandatsreferenz": "MREF", "Glaeubiger-ID": "GLID",
           "Analyse-Vertrags-ID": "VID", "Notiz": "NOTIZ"}
TEXT_COLS = ["Verwendungszweck", "Analyse-Vertrag", "Tags"]

# IBAN kompakt (DE89370400440532013000) oder in 4er-Gruppen (DE89 3704 0044 ...)
IBAN = (r"[A-Z]{2}\d{2}[A-Z0-9]{11,30}"
        r"|[A-Z]{2}\d{2}(?: (?=[A-Z]*\d)[A-Z0-9]{4}){2,7}(?: (?=[A-Z]*\d)[A-Z0-9]{1,3})?")
IBAN_FULL = re.compile(r"[A-Z]{2}\d{2}[A-Z0-9]{11,30}")
IBAN_OR_NUMBER = re.compile(rf"\b(?P<iban>{IBAN})\b|(?P<num>\d{{5,}})")
IBAN_START = re.compile(r"(?=[A-Z]{2}\d{2})")


def is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def norm(text: str) -> str:
    """Vergleichsschlüssel: Groß/klein und Mehrfach-Leerzeichen egal."""
    return " ".join(text.split()).lower()


def literal_pattern(terms, word_boundaries: bool) -> re.Pattern | None:
    terms = sorted({norm(t) for t in terms}, key=len, reverse=True)
    if not terms:
        return None
    alternatives = "|".join(re.escape(t).replace(r"\ ", r"\s+") for t in terms)
    if word_boundaries:
        alternatives = rf"(?<!\w)(?:{alternatives})(?!\w)"
    return re.compile(alternatives, re.IGNORECASE)


def read_list(path: Path) -> list[str]:
    if not path.exists():
        return []
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    return [line.strip() for line in lines if line.strip() and not line.lstrip().startswith("#")]


def load_key() -> bytes:
    if not KEY_FILE.exists():
        KEY_FILE.write_text(secrets.token_hex(32), encoding="utf-8")
        print(f"Neuer Schlüssel -> {KEY_FILE} (nicht teilen!)")
    return bytes.fromhex(KEY_FILE.read_text(encoding="utf-8").strip())


class Anonymizer:
    def __init__(self, key: bytes, keep: set[str], parties: set[str]):
        self.key = key
        self.keep = keep
        hidden = {norm(p) for p in parties} - keep
        self.names = literal_pattern([p for p in hidden if len(p) >= 4], word_boundaries=True)

    def _digest(self, kind: str, value: str) -> bytes:
        return hmac.new(self.key, f"{kind}:{value}".encode(), hashlib.sha256).digest()

    def letters(self, kind: str, value: str, n: int = 6) -> str:
        return "".join(chr(65 + b % 26) for b in self._digest(kind, value)[:n])

    def digits(self, kind: str, value: str, n: int) -> str:
        return "".join(str(b % 10) for b in (self._digest(kind, value) * 2)[:n])

    def party(self, name: str) -> str:
        key = norm(name)
        return name if key in self.keep else f"PARTEI_{self.letters('party', key)}"

    def account_name(self, name: str) -> str:
        key = norm(name)
        return name if key in self.keep else f"KONTO_{self.letters('account', key)}"

    def iban(self, value: str) -> str:
        compact = value.replace(" ", "").upper()
        if not IBAN_FULL.fullmatch(compact):  # z. B. PayPal-Konto ohne IBAN
            return f"ID_{self.letters('account-id', compact, 8)}"
        return compact[:2] + "00" + self.digits("iban", compact, len(compact) - 4)

    def ident(self, prefix: str, value: str) -> str:
        return f"{prefix}_{self.letters(prefix, norm(value), 8)}"

    def text(self, text: str) -> str:
        if self.names:
            text = self.names.sub(lambda m: f"PARTEI_{self.letters('party', norm(m.group(0)))}", text)

        def replace(m: re.Match) -> str:
            if m.group("iban"):
                return self.iban(m.group("iban"))
            return self.digits("num", m.group("num"), len(m.group("num")))

        return IBAN_OR_NUMBER.sub(replace, text)


def contains_iban(text: str, ibans: set[str], lengths: set[int]) -> bool:
    compact = text.replace(" ", "").upper()
    return any(compact[m.start():m.start() + n] in ibans
               for m in IBAN_START.finditer(compact) for n in lengths)


def main(src: Path, keep_file: Path, redact_file: Path) -> None:
    if not keep_file.exists():
        sys.exit(f"{keep_file} fehlt – erst inspect_export.py laufen lassen.")
    keep = {norm(n) for n in read_list(keep_file)}
    redact = literal_pattern(read_list(redact_file), word_boundaries=False)
    if not redact:
        print(f"Hinweis: {redact_file} fehlt/leer – Name/Adresse werden nicht gezielt entfernt.")

    wb = load_workbook(src)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    rows = list(ws.iter_rows(min_row=2))

    parties = {str(r[col[PARTY]].value) for r in rows if not is_empty(r[col[PARTY]].value)}
    original_ibans = {str(r[col[c]].value).replace(" ", "").upper()
                      for r in rows for c in IBAN_COLS if not is_empty(r[col[c]].value)}
    original_ibans = {i for i in original_ibans if IBAN_FULL.fullmatch(i)}
    iban_lengths = {len(i) for i in original_ibans}

    anon = Anonymizer(load_key(), keep, parties)
    rules = {PARTY: anon.party, ACCOUNT_NAME: anon.account_name}
    rules |= {c: anon.iban for c in IBAN_COLS}
    rules |= {c: (lambda v, p=p: anon.ident(p, v)) for c, p in ID_COLS.items()}
    rules |= {c: anon.text for c in TEXT_COLS}

    changed, leaks = Counter(), {}
    for row in rows:
        for name, cell in zip(header, row):
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            new = rules[name](cell.value) if name in rules else cell.value
            if redact:
                new = redact.sub("XXX", new)
            if new != cell.value:
                cell.value = new
                changed[name] += 1
            if (anon.names and anon.names.search(new)) or contains_iban(new, original_ibans, iban_lengths):
                leaks.setdefault(name, []).append(cell.row)

    out = src.with_name(f"{src.stem}_anon.xlsx")
    wb.save(out)

    print(f"Gegenparteien: {len(parties)}, davon Klartext laut keep.txt: "
          f"{sum(norm(p) in keep for p in parties)}")
    print("Geänderte Zellen pro Spalte:")
    for name in header:
        if changed[name]:
            print(f"  {name:<40} {changed[name]:>6}")
    if leaks:
        print("WARNUNG – mögliche Lecks (versteckter Name oder echte IBAN noch enthalten):")
        for name, row_numbers in leaks.items():
            preview = ", ".join(map(str, row_numbers[:10])) + (" …" if len(row_numbers) > 10 else "")
            print(f"  {name:<40} {len(row_numbers):>6}  Zeilen: {preview}")
    else:
        print("Leak-Check: OK")
    print(f"-> {out}")


if __name__ == "__main__":
    if not 2 <= len(sys.argv) <= 4:
        sys.exit("Aufruf: python anonymize.py <export.xlsx> [keep.txt] [redact.txt]")
    args = [Path(a) for a in sys.argv[1:]] + [Path("keep.txt"), Path("redact.txt")][len(sys.argv) - 2:]
    main(*args)
